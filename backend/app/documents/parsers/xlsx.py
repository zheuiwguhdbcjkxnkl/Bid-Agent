from __future__ import annotations

import hashlib
from io import BytesIO

from openpyxl import load_workbook

from app.core.errors import DomainError
from app.documents.ir import DocumentIR, DocumentIRPage, DocumentIRSegment


def parse_xlsx(content: bytes) -> DocumentIR:
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=False,
        )
    except Exception as exc:
        raise DomainError(422, "XLSX_PARSE_FAILED", "XLSX 文件无法解析") from exc

    pages: list[DocumentIRPage] = []
    segments: list[DocumentIRSegment] = []
    for page_no, sheet in enumerate(workbook.worksheets, start=1):
        pages.append(
            DocumentIRPage(
                page_no=page_no,
                metadata_json={"sheet_name": sheet.title},
            )
        )
        cells: list[str] = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells.append(f"{cell.coordinate}\t{cell.value}")
        if not cells:
            continue
        text = "\n".join(cells)
        segments.append(
            DocumentIRSegment(
                segment_type="TABLE",
                page_no=page_no,
                section_path=sheet.title,
                content_text=text,
                source_hash=hashlib.sha256(text.encode("utf-8")).hexdigest(),
                metadata_json={
                    "sheet_name": sheet.title,
                    "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                },
            )
        )
    workbook.close()
    if not segments:
        raise DomainError(422, "XLSX_PARSE_EMPTY", "XLSX 文件没有可解析内容")
    return DocumentIR(
        parser_name="openpyxl",
        parser_version="3.1.5",
        backend_mode="native",
        pages=pages,
        segments=segments,
    )
