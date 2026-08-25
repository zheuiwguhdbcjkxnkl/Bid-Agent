from __future__ import annotations

from io import BytesIO

from docx import Document

from app.documents.parsers.docx import parse_docx


def _docx_content() -> bytes:
    document = Document()
    document.add_heading("第一章 项目概况", level=1)
    document.add_paragraph("本项目建设统一平台。")
    document.add_paragraph("第一项", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "资格"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "资质证书"
    table.cell(1, 1).text = "有效"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_docx_parser_preserves_heading_paragraph_list_and_table_order() -> None:
    ir = parse_docx(_docx_content())

    assert ir.parser_name == "python-docx"
    assert [segment.segment_type for segment in ir.segments] == [
        "HEADING",
        "PARAGRAPH",
        "LIST",
        "TABLE",
    ]
    assert ir.segments[0].section_path == "第一章 项目概况"
    assert "资格\t要求" in ir.segments[3].content_text
    assert ir.segments[3].metadata_json == {"rows": 2, "columns": 2}
